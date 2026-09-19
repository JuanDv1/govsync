"""Pruebas de resolver_hojas y leer() en lectores/ejecucion.py.

TARJETAS: [HU-03][BE-01] (resolver_hojas), [HU-03][BE-04]/[BE-05] (leer())
CUBRE: HU-03 / CA-2 — ambas pestañas (Ejecución y Contratación) se localizan
de forma independiente: la ausencia de una no impide resolver la otra, y no
se exige subirlas por separado (es un solo archivo, dos hojas).
"""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.modules.ingesta.persistence.lectores.ejecucion import (
    LectorEjecucion,
    _leer_contratos_y_registros,
    _leer_rubros,
    resolver_hojas,
)
from app.shared.errors import ArchivoInvalido
from tests.fabricas import (
    COD_B,
    HOJA_CONTRATACION,
    HOJA_EJECUCION,
    construir_ejecucion,
    construir_pdt,
)


def test_resuelve_ambas_hojas_cuando_las_dos_existen() -> None:
    hoja_ejecucion, hoja_contratacion = resolver_hojas(construir_ejecucion(), "presupuestal.xlsx")

    assert hoja_ejecucion == HOJA_EJECUCION
    assert hoja_contratacion == HOJA_CONTRATACION


def test_resuelve_ejecucion_aunque_falte_contratacion() -> None:
    contenido = construir_ejecucion(incluir_contratacion=False)

    hoja_ejecucion, hoja_contratacion = resolver_hojas(contenido, "presupuestal.xlsx")

    assert hoja_ejecucion == HOJA_EJECUCION
    assert hoja_contratacion is None


def test_resuelve_contratacion_aunque_falte_ejecucion() -> None:
    contenido = construir_ejecucion(incluir_ejecucion=False)

    hoja_ejecucion, hoja_contratacion = resolver_hojas(contenido, "presupuestal.xlsx")

    assert hoja_ejecucion is None
    assert hoja_contratacion == HOJA_CONTRATACION


def test_ninguna_pestana_no_lanza_excepcion() -> None:
    contenido = construir_ejecucion(incluir_ejecucion=False, incluir_contratacion=False)

    assert resolver_hojas(contenido, "presupuestal.xlsx") == (None, None)


class TestLeer:
    """[HU-03][BE-04]/[BE-05]: orquesta resolver_hojas (ya probado por
    separado) y cubre CA-4 (pestaña faltante, nombrada), CA-5 (archivo que
    no corresponde) y CA-6 (número y descripción del contrato para la
    matriz)."""

    def test_caso_feliz_devuelve_ejecucion_y_contratacion(self) -> None:
        resultado = LectorEjecucion().leer(
            construir_ejecucion(), "presupuestal.xlsx", vigencia=2026
        )

        assert resultado.conteos == {
            "ejecucion": 1,
            "contratacion": 1,
            "rubros": 1,
            "contratos": 1,
            "registros": 1,
        }
        assert resultado.advertencias == []
        assert resultado.filas["ejecucion"][0]["cod_indicador_producto"] == COD_B
        assert resultado.filas["contratacion"][0]["cod_indicador_producto"] == COD_B

    def test_ca6_aporta_numero_y_descripcion_del_contrato(self) -> None:
        resultado = LectorEjecucion().leer(
            construir_ejecucion(), "presupuestal.xlsx", vigencia=2026
        )

        contrato = resultado.filas["contratacion"][0]
        assert contrato["numero_contrato"] == "C-001"
        assert contrato["descripcion_contrato"] == "Mantenimiento de vías terciarias"

    def test_ca4_rechaza_y_nombra_ejecucion_cuando_falta(self) -> None:
        contenido = construir_ejecucion(incluir_ejecucion=False)

        with pytest.raises(ArchivoInvalido) as exc:
            LectorEjecucion().leer(contenido, "presupuestal.xlsx", vigencia=2026)

        assert exc.value.detalles == {"motivo": "pestana_faltante", "pestana_faltante": "EJECUCION"}

    def test_ca4_rechaza_y_nombra_contratacion_cuando_falta(self) -> None:
        contenido = construir_ejecucion(incluir_contratacion=False)

        with pytest.raises(ArchivoInvalido) as exc:
            LectorEjecucion().leer(contenido, "presupuestal.xlsx", vigencia=2026)

        assert exc.value.detalles == {
            "motivo": "pestana_faltante",
            "pestana_faltante": "CONTRATACION",
        }

    def test_ca5_rechaza_un_archivo_que_no_es_el_presupuestal(self) -> None:
        """Subir el PDT donde se espera el archivo presupuestal: ninguna de
        las dos pestañas aparece, así que es CA-5 (archivo incorrecto), no
        CA-4 (pestaña faltante)."""
        with pytest.raises(ArchivoInvalido) as exc:
            LectorEjecucion().leer(construir_pdt(), "pdt.xlsx", vigencia=2026)

        assert exc.value.detalles["motivo"] == "archivo_no_corresponde"

    def test_fila_con_codigo_de_indicador_invalido_se_descarta_con_advertencia(self) -> None:
        libro = Workbook()
        libro.remove(libro.active)
        hoja = libro.create_sheet(HOJA_EJECUCION)
        hoja.append(["CodigoIndicadorCcpet"])
        hoja.append(["12345"])  # 5 dígitos: no es normalizable a 9
        hoja_contratacion = libro.create_sheet(HOJA_CONTRATACION)
        hoja_contratacion.append(["Cod Indicador Ccpet"])
        hoja_contratacion.append([COD_B])
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorEjecucion().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        assert resultado.filas["ejecucion"] == []
        assert resultado.conteos["ejecucion"] == 0
        assert "12345" in resultado.advertencias[0]

    def test_columnas_opcionales_ausentes_no_rechazan_el_archivo(self) -> None:
        """numero_contrato/descripcion_contrato no son obligatorias (solo CA-6)."""
        libro = Workbook()
        libro.remove(libro.active)
        libro.create_sheet(HOJA_EJECUCION).append(["CodigoIndicadorCcpet"])
        libro[HOJA_EJECUCION].append([COD_B])
        hoja_contratacion = libro.create_sheet(HOJA_CONTRATACION)
        hoja_contratacion.append(["Cod Indicador Ccpet"])
        hoja_contratacion.append([COD_B])
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorEjecucion().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        contrato = resultado.filas["contratacion"][0]
        assert contrato["numero_contrato"] is None
        assert contrato["descripcion_contrato"] is None


def _libro(hoja: str, filas: list[list]) -> bytes:
    """Construye un libro con una sola hoja llamada `hoja` y esas filas
    (la primera es el encabezado) — atajo para los casos de _leer_rubros y
    _leer_contratos_y_registros, que no necesitan la otra pestaña."""
    libro = Workbook()
    libro.remove(libro.active)
    ws = libro.create_sheet(hoja)
    for fila in filas:
        ws.append(fila)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


class TestLeerRubros:
    """[HU-03][BE-06]: extracción de la pestaña de ejecución completa para
    persistencia de Rubro (independiente de _leer_pestana/CA-6)."""

    def test_fila_totales_se_descarta_sin_advertencia(self) -> None:
        contenido = _libro(
            HOJA_EJECUCION,
            [
                ["CodigoRubroNivel", "UltimoNivel"],
                ["1.2.3", "True"],
                ["", "False"],  # fila TOTALES: sin codigo_rubro_nivel
            ],
        )

        rubros, advertencias = _leer_rubros(contenido, "presupuestal.xlsx", HOJA_EJECUCION)

        assert len(rubros) == 1
        assert rubros[0]["codigo_rubro_nivel"] == "1.2.3"
        assert advertencias == []

    def test_codigo_rubro_nivel_duplicado_se_descarta_con_advertencia(self) -> None:
        contenido = _libro(
            HOJA_EJECUCION,
            [
                ["CodigoRubroNivel", "UltimoNivel"],
                ["1.2.3", "True"],
                ["1.2.3", "False"],
            ],
        )

        rubros, advertencias = _leer_rubros(contenido, "presupuestal.xlsx", HOJA_EJECUCION)

        assert len(rubros) == 1
        assert rubros[0]["ultimo_nivel"] is True  # se conserva la primera aparición
        assert len(advertencias) == 1
        assert "1.2.3" in advertencias[0]

    def test_ultimo_nivel_no_reconocible_se_descarta_con_advertencia(self) -> None:
        contenido = _libro(
            HOJA_EJECUCION,
            [
                ["CodigoRubroNivel", "UltimoNivel"],
                ["1.2.3", "N/A"],
            ],
        )

        rubros, advertencias = _leer_rubros(contenido, "presupuestal.xlsx", HOJA_EJECUCION)

        assert rubros == []
        assert len(advertencias) == 1
        assert "UltimoNivel" in advertencias[0]

    def test_codigo_rubro_completo_replica_codigo_rubro_nivel(self) -> None:
        """DECISIÓN TÉCNICA (docstring del módulo): no hay un segundo dato
        distinto para codigo_rubro_completo, se llena con el mismo valor."""
        contenido = _libro(
            HOJA_EJECUCION,
            [
                ["CodigoRubroNivel", "UltimoNivel"],
                ["1.2.3", "True"],
            ],
        )

        rubros, _ = _leer_rubros(contenido, "presupuestal.xlsx", HOJA_EJECUCION)

        assert rubros[0]["codigo_rubro_completo"] == rubros[0]["codigo_rubro_nivel"] == "1.2.3"

    def test_montos_opcionales_se_parsean_con_numero(self) -> None:
        contenido = _libro(
            HOJA_EJECUCION,
            [
                ["CodigoRubroNivel", "UltimoNivel", "ApropiacionDefinitiva"],
                ["1.2.3", "True", "$ 1.218.264.452"],
            ],
        )

        rubros, _ = _leer_rubros(contenido, "presupuestal.xlsx", HOJA_EJECUCION)

        assert rubros[0]["apropiacion_definitiva"] == Decimal("1218264452")

    def test_columna_ancla_ausente_degrada_en_vez_de_rechazar_archivo(self) -> None:
        """Ninguna CA de HU-03 exige rechazar TODO el archivo si falta
        CodigoRubroNivel: se degrada con advertencia (no ArchivoInvalido)."""
        contenido = _libro(HOJA_EJECUCION, [["OtraColumna"], ["x"]])

        rubros, advertencias = _leer_rubros(contenido, "presupuestal.xlsx", HOJA_EJECUCION)

        assert rubros == []
        assert len(advertencias) == 1
        assert "CodigoRubroNivel" in advertencias[0]


class TestLeerContratosYRegistros:
    """[HU-03][BE-06]: extracción de CONTRATACION completa para persistencia
    de Contrato + RegistroPresupuestal (independiente de _leer_pestana/CA-6)."""

    def test_numero_contrato_vacio_se_descarta_con_advertencia(self) -> None:
        # La fila necesita OTRO dato además de NumeroContrato vacío: si toda
        # la fila fuera NaN, `_comun.leer_hoja` ya la descarta (dropna) antes
        # de que esta función la vea — no sería un caso real de este código.
        contenido = _libro(
            HOJA_CONTRATACION,
            [
                ["NumeroContrato", "Objeto"],
                ["", "Contrato sin número"],
            ],
        )

        contratos, registros, advertencias = _leer_contratos_y_registros(
            contenido, "presupuestal.xlsx", HOJA_CONTRATACION
        )

        assert contratos == []
        assert registros == []
        assert len(advertencias) == 1
        assert "NumeroContrato" in advertencias[0]

    def test_agrupa_por_numero_contrato_y_toma_el_maximo_de_valor_contrato_y_pagado(self) -> None:
        """45 filas reales repiten (NumeroContrato, Numero Registro) con
        rubros distintos (docstring del módulo): un Contrato por número,
        un RegistroPresupuestal por fila cruda."""
        contenido = _libro(
            HOJA_CONTRATACION,
            [
                ["NumeroContrato", "Valor Contrato", "Pagos", "Numero Registro"],
                ["C-001", "100000000", "50000000", "REG-1"],
                ["C-001", "80000000", "70000000", "REG-2"],
            ],
        )

        contratos, registros, advertencias = _leer_contratos_y_registros(
            contenido, "presupuestal.xlsx", HOJA_CONTRATACION
        )

        assert advertencias == []
        assert len(contratos) == 1
        assert contratos[0]["numero_contrato"] == "C-001"
        assert contratos[0]["valor_contrato"] == Decimal("100000000")
        assert contratos[0]["valor_pagado"] == Decimal("70000000")
        assert len(registros) == 2
        assert {r["numero_registro"] for r in registros} == {"REG-1", "REG-2"}

    def test_bpin_y_cod_indicador_toman_el_primer_valor_no_nulo_del_grupo(self) -> None:
        contenido = _libro(
            HOJA_CONTRATACION,
            [
                ["NumeroContrato", "Codigo Bpin"],
                ["C-001", ""],
                ["C-001", "2026760010123"],
            ],
        )

        contratos, _, _ = _leer_contratos_y_registros(
            contenido, "presupuestal.xlsx", HOJA_CONTRATACION
        )

        assert contratos[0]["bpin"] == "2026760010123"

    def test_columna_ancla_ausente_degrada_en_vez_de_rechazar_archivo(self) -> None:
        """Ninguna CA de HU-03 exige rechazar TODO el archivo si falta
        NumeroContrato: se degrada con advertencia (no ArchivoInvalido)."""
        contenido = _libro(HOJA_CONTRATACION, [["OtraColumna"], ["x"]])

        contratos, registros, advertencias = _leer_contratos_y_registros(
            contenido, "presupuestal.xlsx", HOJA_CONTRATACION
        )

        assert contratos == []
        assert registros == []
        assert len(advertencias) == 1
        assert "NumeroContrato" in advertencias[0]
