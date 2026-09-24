"""Pruebas de pdt.py, incluyendo `leer()` ([HU-02][BE-03]).

TARJETAS: [HU-02][BE-01] (resolver_hoja_pdt, celdas combinadas en encabezado)
          [HU-02][BE-02] (OBLIGATORIAS)
          [HU-02][BE-03] (leer(): orquesta lo anterior + CA-4 y CA-6)
CUBRE: HU-02 / CA-2, CA-3, CA-4, CA-6.

Dos huecos encontrados al comparar el texto COMPLETO de la tarjeta de Trello
de [HU-02][BE-01] contra lo entregado (no solo su docstring-resumen), y
cerrados aquí:
  1. "Si la pestaña no existe, lanza excepción identificando el problema" —
     no estaba implementado (`_comun.resolver_hoja` deliberadamente devuelve
     `None`, para permitir búsquedas independientes en HU-03). Ver
     `TestResolverHojaPdt`.
  2. "Celdas combinadas en el encabezado" — sin probar. Ver
     `test_leer_hoja_con_celda_combinada_en_el_encabezado`.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.modules.ingesta.persistence.lectores import _comun
from app.modules.ingesta.persistence.lectores.pdt import (
    OBLIGATORIAS,
    LectorPDT,
    alias_columna_programacion,
    resolver_hoja_pdt,
)
from app.shared.errors import ArchivoInvalido
from tests.fabricas import HOJA_PDT, construir_ejecucion, construir_pdt

_ENCABEZADOS_REALES = (
    "Código de indicador de producto (MGA)",
    "Código de indicador de producto (SisPT)",
    "Producto (MGA)",
    "Indicador de Producto(MGA)",
    "Principal",
    "Programación del producto bien o servicio 2026",
    "Total 2026",
)


class TestResolverHojaPdt:
    def test_encuentra_la_pestana_cuando_existe(self) -> None:
        assert resolver_hoja_pdt(construir_pdt(), "pdt.xlsx") == HOJA_PDT

    def test_lanza_archivo_invalido_si_no_existe(self) -> None:
        libro = Workbook()
        libro.active.title = "Otra cosa"
        buffer = io.BytesIO()
        libro.save(buffer)

        with pytest.raises(ArchivoInvalido) as exc:
            resolver_hoja_pdt(buffer.getvalue(), "ejecucion.xlsx")
        assert exc.value.detalles["motivo"] == "hoja_no_encontrada"


def test_leer_hoja_con_celda_combinada_en_el_encabezado() -> None:
    """Una columna de encabezado combinada no debe tumbar ni desplazar a las
    demás: pandas la renombra a "Unnamed: N" y sigue leyendo el resto bien.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(["Código de indicador de producto (MGA)", "Programación", None, "Principal"])
    hoja.merge_cells("B1:C1")
    hoja.append(["040110500", "10", None, "Sí"])
    buffer = io.BytesIO()
    libro.save(buffer)

    df = _comun.leer_hoja(buffer.getvalue(), "Datos", fila_encabezado=0)

    assert df["codigo de indicador de producto (mga)"].tolist() == ["040110500"]
    assert df["principal"].tolist() == ["Sí"]


def test_alias_de_obligatorias_coinciden_con_encabezados_reales() -> None:
    normalizados_reales = {_comun.normalizar_encabezado(c) for c in _ENCABEZADOS_REALES}
    for alias in OBLIGATORIAS.values():
        assert any(_comun.normalizar_encabezado(a) in normalizados_reales for a in alias)


def test_alias_columna_programacion_coincide_con_el_encabezado_real_de_esa_vigencia() -> None:
    generado = alias_columna_programacion(2026)
    normalizados_reales = {_comun.normalizar_encabezado(c) for c in _ENCABEZADOS_REALES}
    assert _comun.normalizar_encabezado(generado) in normalizados_reales


def test_alias_columna_programacion_cambia_con_la_vigencia() -> None:
    assert alias_columna_programacion(2025) != alias_columna_programacion(2026)


class TestLeer:
    """[HU-02][BE-03]: orquesta CA-2/CA-3 (ya probadas por separado) y cubre
    CA-4 (archivo que no corresponde) y CA-6 (código + nombre para la matriz)."""

    def test_caso_feliz_devuelve_las_metas_reconocidas(self) -> None:
        resultado = LectorPDT().leer(construir_pdt(), "pdt.xlsx", vigencia=2026)

        assert resultado.conteos == {"metas": 2}
        assert resultado.advertencias == []
        codigos = {m["cod_indicador_producto"] for m in resultado.filas["metas"]}
        assert codigos == {"040110500", "170202300"}

    def test_ca6_aporta_codigo_y_nombre_del_producto(self) -> None:
        resultado = LectorPDT().leer(construir_pdt(), "pdt.xlsx", vigencia=2026)

        meta = next(
            m for m in resultado.filas["metas"] if m["cod_indicador_producto"] == "040110500"
        )
        assert meta["nombre_producto"] == "Vías terciarias mantenidas"
        assert meta["unidad_medida"] == "Kilómetros"
        assert meta["principal"] is True
        assert meta["meta_cuatrienio"] == 10

    def test_ca3_rechaza_y_nombra_la_columna_que_falta(self) -> None:
        with pytest.raises(ArchivoInvalido) as exc:
            LectorPDT().leer(
                construir_pdt(incluir_principal=False), "pdt_sin_principal.xlsx", vigencia=2026
            )

        assert exc.value.detalles["motivo"] == "columnas_faltantes"
        assert exc.value.detalles["columnas_faltantes"] == ["Principal"]

    def test_ca3_rechaza_si_la_vigencia_no_tiene_columna_de_programacion(self) -> None:
        """El archivo es de 2026; se pide leerlo para 2025 (columna inexistente)."""
        with pytest.raises(ArchivoInvalido) as exc:
            LectorPDT().leer(construir_pdt(), "pdt.xlsx", vigencia=2025)

        assert exc.value.detalles["motivo"] == "columnas_faltantes"

    def test_ca4_rechaza_un_archivo_que_no_es_el_pdt(self) -> None:
        """Subir el archivo de ejecución donde se espera un PDT."""
        with pytest.raises(ArchivoInvalido) as exc:
            LectorPDT().leer(construir_ejecucion(), "ejecucion.xlsx", vigencia=2026)

        assert exc.value.detalles["motivo"] == "hoja_no_encontrada"

    def test_fila_con_codigo_de_indicador_invalido_se_descarta_con_advertencia(self) -> None:
        libro = Workbook()
        libro.remove(libro.active)
        hoja = libro.create_sheet(HOJA_PDT)
        hoja.append(
            [
                "Código de indicador de producto (MGA)",
                "Principal",
                "Programación del producto bien o servicio 2026",
            ]
        )
        hoja.append(["12345", "Sí", "10"])  # 5 dígitos: no es normalizable a 9
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorPDT().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        assert resultado.filas["metas"] == []
        assert resultado.conteos == {"metas": 0}
        assert "12345" in resultado.advertencias[0]

    def test_fila_con_principal_no_reconocible_se_descarta_con_advertencia(self) -> None:
        libro = Workbook()
        libro.remove(libro.active)
        hoja = libro.create_sheet(HOJA_PDT)
        hoja.append(
            [
                "Código de indicador de producto (MGA)",
                "Principal",
                "Programación del producto bien o servicio 2026",
            ]
        )
        hoja.append(["040110500", "Tal vez", "5"])
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorPDT().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        assert resultado.filas["metas"] == []
        assert "no reconocible" in resultado.advertencias[0]

    def test_fila_con_meta_no_numerica_no_se_descarta_solo_advierte(self) -> None:
        """A diferencia de código/principal, una meta no numérica no invalida
        la fila entera: la meta queda en None y se advierte (CA-6 solo exige
        código y nombre; el valor de meta es adicional)."""
        libro = Workbook()
        libro.remove(libro.active)
        hoja = libro.create_sheet(HOJA_PDT)
        hoja.append(
            [
                "Código de indicador de producto (MGA)",
                "Principal",
                "Programación del producto bien o servicio 2026",
            ]
        )
        hoja.append(["040110500", "No", "sin dato"])
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorPDT().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        assert len(resultado.filas["metas"]) == 1
        assert resultado.filas["metas"][0]["meta_cuatrienio"] is None
        assert len(resultado.advertencias) == 1

    def test_columnas_opcionales_ausentes_no_rechazan_el_archivo(self) -> None:
        """nombre_producto/unidad_medida no son obligatorias (solo CA-6)."""
        libro = Workbook()
        libro.remove(libro.active)
        hoja = libro.create_sheet(HOJA_PDT)
        hoja.append(
            [
                "Código de indicador de producto (MGA)",
                "Principal",
                "Programación del producto bien o servicio 2026",
            ]
        )
        hoja.append(["040110500", "Sí", "10"])
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorPDT().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        meta = resultado.filas["metas"][0]
        assert meta["nombre_producto"] is None
        assert meta["unidad_medida"] is None

    def test_columnas_agregadas_2026_09_23_se_extraen_cuando_estan_presentes(self) -> None:
        """entidad_territorial/nombre_plan/fecha_creacion_plan (metadato del
        plan) + linea_estrategica/sector/programa/ods/tipo_acumulacion (por
        meta) — ver docs/DATOS.md."""
        libro = Workbook()
        libro.remove(libro.active)
        hoja = libro.create_sheet(HOJA_PDT)
        hoja.append(
            [
                "Código de indicador de producto (MGA)",
                "Principal",
                "Programación del producto bien o servicio 2026",
                "Código de indicador de producto (SisPT)",
                "Código de producto (MGA)",
                "Entidad Territorial",
                "Nombre del Plan",
                "Fecha de creación del plan",
                "Línea estratégica",
                "Código del sector (MGA)",
                "Sector (MGA)",
                "Código del programa (MGA)",
                "Programa (MGA)",
                "Código ODS",
                "ODS",
                "Tipo de acumulación",
            ]
        )
        hoja.append(
            [
                "040110500",
                "Sí",
                "10",
                "IP-63",
                "0401",
                "Santa Rosa, Cauca",
                "Plan de Desarrollo Municipal 2024-2027",
                "2024-01-15",
                "Vías para la competitividad",
                "04",
                "Transporte",
                "0401",
                "Infraestructura vial",
                "09",
                "Industria, innovación e infraestructura",
                "Sumable",
            ]
        )
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorPDT().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        meta = resultado.filas["metas"][0]
        assert meta["cod_indicador_sistp"] == "IP-63"
        assert meta["codigo_producto_mga"] == "0401"
        assert meta["entidad_territorial"] == "Santa Rosa, Cauca"
        assert meta["nombre_plan"] == "Plan de Desarrollo Municipal 2024-2027"
        assert meta["fecha_creacion_plan"].isoformat() == "2024-01-15"
        assert meta["linea_estrategica"] == "Vías para la competitividad"
        assert meta["codigo_sector"] == "04"
        assert meta["sector"] == "Transporte"
        assert meta["codigo_programa"] == "0401"
        assert meta["programa"] == "Infraestructura vial"
        assert meta["codigo_ods"] == "09"
        assert meta["ods"] == "Industria, innovación e infraestructura"
        assert meta["tipo_acumulacion"] == "Sumable"
