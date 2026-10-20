"""Pruebas del endpoint del Plan de Acción Consolidado (capa API).

TARJETA: [E04-HU01]
CUBRE: HU01-CA01 (corte actual), CA04 (solo lectura -- por ausencia de
verbos de escritura), CA09 (estado vacío -- 404 cuando no hay corte
registrado), CA10 (restricción de rol -- D26, PENDIENTE DE RATIFICAR).

Mismo patrón que `test_router_trazabilidad.py`: sin dobles, SQL real contra
SQLite en memoria vía `cliente`/`sesion` de conftest.py. El rol por defecto
del fixture `cliente` es GESTOR (ver conftest.py) -- como D26 NO le da
acceso a este endpoint, cada prueba que SÍ espera 200 sobreescribe
`obtener_usuario_actual` con un usuario administrador o visitante (mismo
mecanismo que `test_autorizacion_rutas.py::TestRolVisitante`).
"""

from __future__ import annotations

import uuid
from datetime import date
from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.dependencias import obtener_usuario_actual
from app.modules.cortes.domain.entidades import EstadoCorte, TipoArchivoFuente
from app.modules.cortes.persistence.models import (
    ArchivoFuenteORM,
    ContratoORM,
    CorteORM,
    MetaORM,
    ProyectoIndicadorORM,
    ProyectoORM,
    RegistroPresupuestalORM,
    RubroORM,
)
from app.modules.identidad.domain.entidades import Rol, Usuario

_ADMINISTRADOR = Usuario(
    id="a1", correo="admin@santarosa.gov.co", roles=frozenset({Rol.ADMINISTRADOR})
)
_VISITANTE = Usuario(id="v1", correo="visitante@santarosa.gov.co", roles=frozenset({Rol.VISITANTE}))
_GESTOR = Usuario(id="g1", correo="gestor@santarosa.gov.co", roles=frozenset({Rol.GESTOR}))


def _crear_corte(
    sesion,
    *,
    vigencia: int = 2026,
    fecha: date = date(2026, 1, 1),
    estado: EstadoCorte = EstadoCorte.BORRADOR,
) -> uuid.UUID:
    corte_id = uuid.uuid4()
    sesion.add(CorteORM(id=corte_id, vigencia=vigencia, fecha_corte=fecha, estado=estado))
    sesion.flush()
    return corte_id


def _marcar_archivo_cargado(sesion, corte_id: uuid.UUID, tipo: TipoArchivoFuente) -> None:
    sesion.add(
        ArchivoFuenteORM(
            id=uuid.uuid4(),
            corte_id=corte_id,
            tipo=tipo,
            nombre_archivo=f"{tipo.value.lower()}.xlsx",
        )
    )


class TestRestriccionDeRolCA10:
    """D26: administrador + visitante tienen acceso; gestor no (ver
    docstring de plan_accion/api/router.py)."""

    def test_sin_token_devuelve_401(self, sesion) -> None:
        from fastapi.testclient import TestClient

        from app.core.database import get_session
        from app.main import crear_app

        def _sesion_de_prueba():
            yield sesion

        app = crear_app()
        app.dependency_overrides[get_session] = _sesion_de_prueba
        cliente_sin_usuario = TestClient(app)

        respuesta = cliente_sin_usuario.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 401

    def test_gestor_no_tiene_acceso(self, cliente) -> None:
        # El fixture `cliente` ya usa GESTOR por defecto (conftest.py).
        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 403
        cuerpo = respuesta.json()
        assert cuerpo["codigo"] == "permiso_insuficiente"
        assert set(cuerpo["detalles"]["roles_requeridos"]) == {"administrador", "visitante"}

    def test_administrador_tiene_acceso(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 200

    def test_visitante_tiene_acceso(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _VISITANTE
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 200


class TestEstadoVacioCA09:
    def test_sin_ningun_corte_registrado_devuelve_404(self, cliente) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 404
        assert respuesta.json()["codigo"] == "recurso_no_encontrado"

    def test_con_un_corte_solo_borrador_devuelve_404(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        _crear_corte(sesion, estado=EstadoCorte.BORRADOR)
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 404


class TestFuentesIncompletasCA01:
    def test_devuelve_409_si_falta_una_fuente(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        _marcar_archivo_cargado(sesion, corte_id, TipoArchivoFuente.PDT)
        _marcar_archivo_cargado(sesion, corte_id, TipoArchivoFuente.EJECUCION)
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 409
        assert respuesta.json()["detalles"]["archivos_faltantes"] == ["PROYECTOS"]


class TestPlanAccionActualConDatos:
    """CA01/CA02/CA04/CA05: estructura de la respuesta con datos reales."""

    def test_devuelve_la_meta_agregada_con_su_resumen(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)
        sesion.add(
            MetaORM(
                id=uuid.uuid4(),
                corte_id=corte_id,
                cod_indicador_producto="040110500",
                nombre_producto="Vías pavimentadas",
                sector="Infraestructura",
                es_principal=True,
            )
        )
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 200
        cuerpo = respuesta.json()
        assert cuerpo["corte_id"] == str(corte_id)
        assert len(cuerpo["filas"]) == 1
        fila = cuerpo["filas"][0]
        assert fila["cod_indicador_producto"] == "040110500"
        assert fila["sector"] == "Infraestructura"
        # CA04: ningún campo de la fila sugiere edición -- solo lectura.
        assert fila["avance_fisico_pct"] is None
        # Sin rubro/BPIN/contrato: CA05 la clasifica "sin_informacion".
        assert fila["estado_cruce"] == "sin_informacion"
        assert cuerpo["resumen"] == {
            "total": 1,
            "completo": 0,
            "parcial": 0,
            "sin_informacion": 1,
        }

    def test_corte_registrado_sin_ninguna_meta_devuelve_filas_vacias_y_resumen_en_cero(
        self, cliente, sesion
    ) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 200
        cuerpo = respuesta.json()
        assert cuerpo["filas"] == []
        assert cuerpo["resumen"]["total"] == 0


class TestFiltrosQueryParamsE04HU02:
    """[E04-HU02] filtros vía query params, de punta a punta (SQL real +
    agregación + filtrado), igual patrón que las clases de arriba."""

    def _preparar_corte_con_dos_metas(self, sesion):
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)
        sesion.add(
            MetaORM(
                id=uuid.uuid4(),
                corte_id=corte_id,
                cod_indicador_producto="040110500",
                nombre_producto="Vías pavimentadas",
                sector="Infraestructura",
                es_principal=True,
            )
        )
        sesion.add(
            MetaORM(
                id=uuid.uuid4(),
                corte_id=corte_id,
                cod_indicador_producto="170202300",
                nombre_producto="Escuelas construidas",
                sector="Educación",
                es_principal=True,
            )
        )
        sesion.commit()
        return corte_id

    def test_filtra_por_sector(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_dos_metas(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/actual", params={"sector": "Educación"})

        assert respuesta.status_code == 200
        filas = respuesta.json()["filas"]
        assert len(filas) == 1
        assert filas[0]["cod_indicador_producto"] == "170202300"
        assert respuesta.json()["resumen"]["total"] == 1

    def test_busqueda_sin_distinguir_tildes(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_dos_metas(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/actual", params={"busqueda": "Vias"})

        assert respuesta.status_code == 200
        filas = respuesta.json()["filas"]
        assert len(filas) == 1
        assert filas[0]["cod_indicador_producto"] == "040110500"

    def test_sin_resultados_devuelve_lista_vacia_y_resumen_en_cero(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_dos_metas(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/actual", params={"sector": "Salud"})

        assert respuesta.status_code == 200
        cuerpo = respuesta.json()
        assert cuerpo["filas"] == []
        assert cuerpo["resumen"]["total"] == 0

    def test_sin_query_params_devuelve_todo_ca08(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_dos_metas(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/actual")

        assert respuesta.status_code == 200
        assert len(respuesta.json()["filas"]) == 2

    def test_sector_repetido_combina_con_or(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_dos_metas(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/actual",
            params=[("sector", "Educación"), ("sector", "Infraestructura")],
        )

        assert respuesta.status_code == 200
        assert len(respuesta.json()["filas"]) == 2

    def test_filtra_por_indicador_exacto_d27(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_dos_metas(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/actual",
            params={"cod_indicador_producto": "040110500"},
        )

        assert respuesta.status_code == 200
        filas = respuesta.json()["filas"]
        assert len(filas) == 1
        assert filas[0]["cod_indicador_producto"] == "040110500"

    def test_indicador_no_hace_match_parcial_a_diferencia_de_busqueda(
        self, cliente, sesion
    ) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_dos_metas(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/actual",
            params={"cod_indicador_producto": "0401"},
        )

        assert respuesta.status_code == 200
        assert respuesta.json()["filas"] == []


class TestExportarExcelE04HU03:
    """[E04-HU03] GET /plan-accion/exportar -- mismo pipeline que /actual
    (ver _obtener_filas_filtradas en el router), salida .xlsx."""

    def _preparar_corte_con_una_meta(self, sesion):
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)
        sesion.add(
            MetaORM(
                id=uuid.uuid4(),
                corte_id=corte_id,
                cod_indicador_producto="040110500",
                nombre_producto="Vías pavimentadas",
                sector="Infraestructura",
                es_principal=True,
            )
        )
        sesion.commit()
        return corte_id

    def test_mismo_control_de_rol_que_actual_ca10(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _GESTOR
        self._preparar_corte_con_una_meta(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/exportar")

        assert respuesta.status_code == 403

    def test_devuelve_xlsx_descargable_ca01(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_una_meta(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/exportar")

        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "Plan_Accion_2026_Corte_2026-01-01.xlsx" in respuesta.headers["content-disposition"]

    def test_contenido_del_archivo_trae_la_meta_ca01(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_una_meta(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/exportar")

        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro["Plan de Acción"]
        assert hoja.cell(row=7, column=1).value == "Sector"
        assert hoja.cell(row=8, column=2).value == "040110500"

    def test_respeta_filtros_igual_que_actual_ca02(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)

        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)

        sesion.add_all(
            [
                MetaORM(
                    id=uuid.uuid4(),
                    corte_id=corte_id,
                    cod_indicador_producto="040110500",
                    nombre_producto="Meta Educación",
                    sector="Educación",
                    es_principal=True,
                ),
                MetaORM(
                    id=uuid.uuid4(),
                    corte_id=corte_id,
                    cod_indicador_producto="170202300",
                    nombre_producto="Meta Infraestructura",
                    sector="Infraestructura",
                    es_principal=True,
                ),
            ]
        )
        sesion.commit()

        respuesta = cliente.get(
            "/api/v1/plan-accion/exportar",
            params={"sector": "Educación"},
        )

        assert respuesta.status_code == 200

        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro["Plan de Acción"]

        assert hoja.cell(row=8, column=1).value == "Educación"
        assert hoja.cell(row=8, column=2).value == "040110500"

        # El segundo registro no debe aparecer después de aplicar el filtro.
        assert hoja.cell(row=9, column=2).value is None

    def test_valores_exportados_coinciden_con_json_ca06(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)

        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)

        sesion.add(
            MetaORM(
                id=uuid.uuid4(),
                corte_id=corte_id,
                cod_indicador_producto="040110500",
                nombre_producto="Vías pavimentadas",
                sector="Educación",
                es_principal=True,
            )
        )
        sesion.commit()

        respuesta_json = cliente.get("/api/v1/plan-accion/actual")
        respuesta_xlsx = cliente.get("/api/v1/plan-accion/exportar")

        assert respuesta_json.status_code == 200
        assert respuesta_xlsx.status_code == 200

        cuerpo = respuesta_json.json()
        fila_json = cuerpo["filas"][0]

        libro = load_workbook(BytesIO(respuesta_xlsx.content), data_only=True)
        hoja = libro["Plan de Acción"]

        fila_xlsx = [hoja.cell(row=8, column=columna).value for columna in range(1, 13)]

        valores_json = [
            fila_json["sector"],
            fila_json["cod_indicador_producto"],
            fila_json["nombre_producto"],
            fila_json["meta_programada"],
            fila_json["cod_bpin"],
            ", ".join(fila_json["rubros"]) if fila_json["rubros"] else None,
            fila_json["presupuesto_definitivo"],
            fila_json["ejecucion"],
            ", ".join(fila_json["contratos"]) if fila_json["contratos"] else None,
            fila_json["avance_fisico_pct"],
            fila_json["avance_financiero_pct"],
            fila_json["estado_cruce"],
        ]

        assert fila_xlsx == valores_json

        assert cuerpo["resumen"]["total"] == 1

    def _preparar_corte_con_datos_completos(self, sesion):
        """Meta 1 con BPIN, 2 rubros, contrato y montos (estado completo);
        meta 2 sin ningún cruce (sin_informacion)."""
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)

        cod = "040110500"
        proyecto_id = uuid.uuid4()
        rubro_1, rubro_2, contrato_id = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
        sesion.add_all(
            [
                MetaORM(
                    id=uuid.uuid4(),
                    corte_id=corte_id,
                    cod_indicador_producto=cod,
                    nombre_producto="Vías pavimentadas",
                    sector="Infraestructura",
                    meta_cuatrienio=Decimal("120"),
                    es_principal=True,
                ),
                MetaORM(
                    id=uuid.uuid4(),
                    corte_id=corte_id,
                    cod_indicador_producto="170202300",
                    nombre_producto="Escuelas construidas",
                    sector="Educación",
                    es_principal=True,
                ),
                ProyectoORM(id=proyecto_id, corte_id=corte_id, bpin="2024000010001"),
            ]
        )
        sesion.flush()
        sesion.add(
            ProyectoIndicadorORM(
                id=uuid.uuid4(), proyecto_id=proyecto_id, cod_indicador_producto=cod
            )
        )
        for rubro_id, nivel, apropiado, obligado in (
            (rubro_1, "R-A", Decimal("1000000"), Decimal("250000")),
            (rubro_2, "R-B", Decimal("3000000"), Decimal("500000")),
        ):
            sesion.add(
                RubroORM(
                    id=rubro_id,
                    corte_id=corte_id,
                    codigo_rubro_nivel=nivel,
                    codigo_sector_ccpet="22",
                    cod_indicador_producto=cod,
                    ultimo_nivel=True,
                    apropiacion_definitiva=apropiado,
                    obligacion_acumulada=obligado,
                )
            )
        sesion.add(
            ContratoORM(
                id=contrato_id,
                corte_id=corte_id,
                numero_contrato="C-001",
                tipo_gasto="INVERSION",
            )
        )
        sesion.flush()
        sesion.add_all(
            [
                RegistroPresupuestalORM(id=uuid.uuid4(), contrato_id=contrato_id, rubro_id=r)
                for r in (rubro_1, rubro_2)
            ]
        )
        sesion.commit()
        return corte_id

    @staticmethod
    def _normalizar(valor):
        """JSON serializa Decimal como texto/float; el XLSX, como número."""
        if isinstance(valor, str):
            try:
                return float(valor)
            except ValueError:
                return valor
        return float(valor) if isinstance(valor, int | float | Decimal) else valor

    def test_cifras_con_datos_completos_coinciden_con_json_ca06(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_datos_completos(sesion)

        filas_json = cliente.get("/api/v1/plan-accion/actual").json()["filas"]
        respuesta = cliente.get("/api/v1/plan-accion/exportar")
        hoja = load_workbook(BytesIO(respuesta.content))["Plan de Acción"]

        assert len(filas_json) == 2
        campos = (
            "sector",
            "cod_indicador_producto",
            "nombre_producto",
            "meta_programada",
            "cod_bpin",
            "rubros",
            "presupuesto_definitivo",
            "ejecucion",
            "contratos",
            "avance_fisico_pct",
            "avance_financiero_pct",
            "estado_cruce",
        )
        for desplazamiento, fila_json in enumerate(filas_json):
            for columna, campo in enumerate(campos, start=1):
                esperado = fila_json[campo]
                if isinstance(esperado, list):
                    esperado = ", ".join(esperado) if esperado else None
                obtenido = hoja.cell(row=8 + desplazamiento, column=columna).value
                assert self._normalizar(obtenido) == self._normalizar(esperado), (
                    f"fila {desplazamiento + 1}, columna {campo}: {obtenido!r} != {esperado!r}"
                )

        # Guardia: la meta completa de verdad trae cifras (no None == None).
        completa = next(f for f in filas_json if f["cod_indicador_producto"] == "040110500")
        assert completa["estado_cruce"] == "completo"
        assert self._normalizar(completa["presupuesto_definitivo"]) == 4000000.0
        assert self._normalizar(completa["ejecucion"]) == 750000.0
        assert self._normalizar(completa["avance_financiero_pct"]) == 18.75
        assert len(completa["rubros"]) == 2

    def test_formatos_numericos_en_archivo_exportado_ca03(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_datos_completos(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/exportar")
        hoja = load_workbook(BytesIO(respuesta.content))["Plan de Acción"]

        fila = next(r for r in hoja.iter_rows(min_row=8) if r[1].value == "040110500")
        assert fila[6].value == 4000000 and fila[6].number_format == '#,##0 "COP"'
        assert fila[7].value == 750000 and fila[7].number_format == '#,##0 "COP"'
        assert fila[10].value == 18.75 and fila[10].number_format == '0.0"%"'

    def test_visitante_puede_exportar(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _VISITANTE
        self._preparar_corte_con_una_meta(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/exportar")

        assert respuesta.status_code == 200
        assert load_workbook(BytesIO(respuesta.content))["Plan de Acción"]["A7"].value == "Sector"

    def test_openapi_documenta_summary_y_respuestas_de_error(self, cliente) -> None:
        operacion = cliente.app.openapi()["paths"]["/api/v1/plan-accion/exportar"]["get"]

        assert operacion["summary"]
        assert {"200", "401", "403", "404", "409"} <= set(operacion["responses"])

    def test_sin_token_devuelve_401(self, sesion) -> None:
        from app.core.database import get_session
        from app.main import crear_app

        def _sesion_de_prueba():
            yield sesion

        app = crear_app()
        app.dependency_overrides[get_session] = _sesion_de_prueba

        respuesta = TestClient(app).get("/api/v1/plan-accion/exportar")

        assert respuesta.status_code == 401

    def test_fuentes_incompletas_devuelve_409(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_id = _crear_corte(sesion, estado=EstadoCorte.REGISTRADO)
        _marcar_archivo_cargado(sesion, corte_id, TipoArchivoFuente.PDT)
        _marcar_archivo_cargado(sesion, corte_id, TipoArchivoFuente.EJECUCION)
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/exportar")

        assert respuesta.status_code == 409
        assert respuesta.json()["detalles"]["archivos_faltantes"] == ["PROYECTOS"]

    def test_filtros_combinados_y_cabecera_los_describen_ca02_ca04(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_datos_completos(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/exportar",
            params={
                "sector": ["Infraestructura", "Educación"],
                "estado_cruce": "completo",
                "cod_indicador_producto": "040110500",
                "busqueda": "Vías",
            },
        )

        assert respuesta.status_code == 200
        hoja = load_workbook(BytesIO(respuesta.content))["Plan de Acción"]
        # AND entre criterios: solo sobrevive la meta completa.
        assert hoja.cell(row=8, column=2).value == "040110500"
        assert hoja.cell(row=9, column=2).value is None
        cabecera = "\n".join(str(hoja.cell(row=i, column=1).value) for i in range(1, 7))
        assert "Sector: Infraestructura, Educación" in cabecera
        assert "Estado de cruce: completo" in cabecera
        assert "Indicador: 040110500" in cabecera
        assert 'Búsqueda: "Vías"' in cabecera
        assert "1 metas" in cabecera
        # El nombre del archivo depende del corte, no de los filtros.
        assert "Plan_Accion_2026_Corte_2026-01-01.xlsx" in respuesta.headers["content-disposition"]

    def test_sin_corte_registrado_devuelve_404_ca09(self, cliente) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR

        respuesta = cliente.get("/api/v1/plan-accion/exportar")

        assert respuesta.status_code == 404

    def test_metadatos_incluyen_el_usuario_que_exporta_ca04(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_una_meta(sesion)

        respuesta = cliente.get("/api/v1/plan-accion/exportar")

        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro["Plan de Acción"]
        texto_cabecera = "\n".join(str(hoja.cell(row=i, column=1).value) for i in range(1, 7))
        assert _ADMINISTRADOR.correo in texto_cabecera

    def test_error_al_generar_excel_devuelve_500_generico_ca08(
        self, cliente, sesion, monkeypatch
    ) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        self._preparar_corte_con_una_meta(sesion)

        def _fallar_al_generar_excel(*args, **kwargs):
            raise RuntimeError("detalle interno que no debe exponerse")

        monkeypatch.setattr(
            "app.modules.plan_accion.api.router.construir_libro_excel",
            _fallar_al_generar_excel,
        )

        cliente_sin_raise = TestClient(
            cliente.app,
            raise_server_exceptions=False,
        )

        respuesta = cliente_sin_raise.get("/api/v1/plan-accion/exportar")

        assert respuesta.status_code == 500

        cuerpo = respuesta.json()
        assert cuerpo["codigo"] == "error_interno"
        assert cuerpo["mensaje"] == "Ocurrió un error inesperado procesando la solicitud."
        assert cuerpo["detalles"] == {}
        assert "detalle interno que no debe exponerse" not in respuesta.text
        assert "Traceback" not in respuesta.text


class TestCortesDisponiblesE04HU04:
    """[E04-HU04] CA01/CA02: GET /plan-accion/cortes-disponibles."""

    def test_mismo_control_de_rol_ca10(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _GESTOR

        respuesta = cliente.get("/api/v1/plan-accion/cortes-disponibles", params={"vigencia": 2026})

        assert respuesta.status_code == 403

    def test_solo_devuelve_cortes_registrado_de_la_vigencia_pedida(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_registrado_2026 = _crear_corte(
            sesion, vigencia=2026, fecha=date(2026, 3, 1), estado=EstadoCorte.REGISTRADO
        )
        _crear_corte(sesion, vigencia=2026, fecha=date(2026, 2, 1), estado=EstadoCorte.BORRADOR)
        _crear_corte(sesion, vigencia=2025, fecha=date(2025, 12, 1), estado=EstadoCorte.REGISTRADO)
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/cortes-disponibles", params={"vigencia": 2026})

        assert respuesta.status_code == 200
        cuerpo = respuesta.json()
        assert len(cuerpo) == 1
        assert cuerpo[0]["id"] == str(corte_registrado_2026)
        assert cuerpo[0]["fecha_corte"] == "2026-03-01"

    def test_ordena_por_fecha_de_corte_ascendente(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_tardio = _crear_corte(
            sesion, vigencia=2026, fecha=date(2026, 9, 1), estado=EstadoCorte.REGISTRADO
        )
        corte_temprano = _crear_corte(
            sesion, vigencia=2026, fecha=date(2026, 1, 1), estado=EstadoCorte.REGISTRADO
        )
        sesion.commit()

        respuesta = cliente.get("/api/v1/plan-accion/cortes-disponibles", params={"vigencia": 2026})

        cuerpo = respuesta.json()
        assert [fila["id"] for fila in cuerpo] == [str(corte_temprano), str(corte_tardio)]

    def test_sin_cortes_de_la_vigencia_devuelve_lista_vacia_ca02(self, cliente) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR

        respuesta = cliente.get("/api/v1/plan-accion/cortes-disponibles", params={"vigencia": 2030})

        assert respuesta.status_code == 200
        assert respuesta.json() == []


class TestCompararE04HU04:
    """[E04-HU04] CA03 a CA08: GET /plan-accion/comparar y
    /plan-accion/comparar/exportar."""

    def _crear_corte_registrado_con_meta(
        self,
        sesion,
        *,
        vigencia: int,
        fecha: date,
        cod_indicador_producto: str,
        obligacion_acumulada,
        nombre_producto: str = "Vías pavimentadas",
        sector: str = "Infraestructura",
    ) -> uuid.UUID:
        corte_id = _crear_corte(
            sesion, vigencia=vigencia, fecha=fecha, estado=EstadoCorte.REGISTRADO
        )
        for tipo in TipoArchivoFuente:
            _marcar_archivo_cargado(sesion, corte_id, tipo)
        sesion.add(
            MetaORM(
                id=uuid.uuid4(),
                corte_id=corte_id,
                cod_indicador_producto=cod_indicador_producto,
                nombre_producto=nombre_producto,
                sector=sector,
                es_principal=True,
            )
        )
        sesion.add(
            RubroORM(
                id=uuid.uuid4(),
                corte_id=corte_id,
                codigo_rubro_nivel=f"{cod_indicador_producto}-RUBRO-1",
                cod_indicador_producto=cod_indicador_producto,
                ultimo_nivel=True,
                codigo_sector_ccpet="05",
                obligacion_acumulada=obligacion_acumulada,
            )
        )
        return corte_id

    def _preparar_dos_cortes(self, sesion):
        """Corte anterior (enero, ejecución 100) y posterior (marzo, ejecución
        150) de la misma vigencia -- la meta "170202300" solo existe en el
        corte posterior (CA06: "solo en Corte B")."""
        corte_a = self._crear_corte_registrado_con_meta(
            sesion,
            vigencia=2026,
            fecha=date(2026, 1, 1),
            cod_indicador_producto="040110500",
            obligacion_acumulada=100,
        )
        corte_b = self._crear_corte_registrado_con_meta(
            sesion,
            vigencia=2026,
            fecha=date(2026, 3, 1),
            cod_indicador_producto="040110500",
            obligacion_acumulada=150,
        )
        sesion.add(
            MetaORM(
                id=uuid.uuid4(),
                corte_id=corte_b,
                cod_indicador_producto="170202300",
                nombre_producto="Escuelas construidas",
                sector="Educación",
                es_principal=True,
            )
        )
        sesion.commit()
        return corte_a, corte_b

    def test_mismo_control_de_rol_ca10(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _GESTOR
        corte_a, corte_b = self._preparar_dos_cortes(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar",
            params={"corte_id_1": str(corte_a), "corte_id_2": str(corte_b)},
        )

        assert respuesta.status_code == 403

    def test_compara_y_calcula_delta_ca04_ca05(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_a, corte_b = self._preparar_dos_cortes(sesion)

        # Query params en orden invertido (B, luego A) -- CA03: el resultado
        # no debe depender del orden de llegada.
        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar",
            params={"corte_id_1": str(corte_b), "corte_id_2": str(corte_a)},
        )

        assert respuesta.status_code == 200
        cuerpo = respuesta.json()
        assert cuerpo["corte_a"]["corte_id"] == str(corte_a)
        assert cuerpo["corte_b"]["corte_id"] == str(corte_b)

        filas = {fila["cod_indicador_producto"]: fila for fila in cuerpo["filas"]}
        assert filas["040110500"]["presencia"] == "ambos"
        assert filas["040110500"]["ejecucion_a"] == "100.00"
        assert filas["040110500"]["ejecucion_b"] == "150.00"
        assert filas["040110500"]["ejecucion_delta"] == "50.00"
        assert filas["170202300"]["presencia"] == "solo_corte_b"
        assert filas["170202300"]["ejecucion_delta"] is None

    def test_solo_cambios_oculta_metas_sin_delta_ca06_ca07(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_a = self._crear_corte_registrado_con_meta(
            sesion,
            vigencia=2026,
            fecha=date(2026, 1, 1),
            cod_indicador_producto="040110500",
            obligacion_acumulada=100,
        )
        corte_b = self._crear_corte_registrado_con_meta(
            sesion,
            vigencia=2026,
            fecha=date(2026, 3, 1),
            cod_indicador_producto="040110500",
            obligacion_acumulada=100,  # sin cambio
        )
        sesion.commit()

        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar",
            params={
                "corte_id_1": str(corte_a),
                "corte_id_2": str(corte_b),
                "solo_cambios": True,
            },
        )

        assert respuesta.status_code == 200
        assert respuesta.json()["filas"] == []

    def test_mismo_corte_dos_veces_devuelve_422_ca03(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_a, _ = self._preparar_dos_cortes(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar",
            params={"corte_id_1": str(corte_a), "corte_id_2": str(corte_a)},
        )

        assert respuesta.status_code == 422
        assert respuesta.json()["codigo"] == "regla_de_negocio_violada"

    def test_vigencias_distintas_devuelve_422_ca03(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_2026 = self._crear_corte_registrado_con_meta(
            sesion,
            vigencia=2026,
            fecha=date(2026, 1, 1),
            cod_indicador_producto="040110500",
            obligacion_acumulada=100,
        )
        corte_2025 = self._crear_corte_registrado_con_meta(
            sesion,
            vigencia=2025,
            fecha=date(2025, 1, 1),
            cod_indicador_producto="040110500",
            obligacion_acumulada=100,
        )
        sesion.commit()

        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar",
            params={"corte_id_1": str(corte_2026), "corte_id_2": str(corte_2025)},
        )

        assert respuesta.status_code == 422
        assert respuesta.json()["detalles"]["motivo"] == "vigencias_distintas"

    def test_corte_en_borrador_devuelve_409_ca07(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_registrado = self._crear_corte_registrado_con_meta(
            sesion,
            vigencia=2026,
            fecha=date(2026, 1, 1),
            cod_indicador_producto="040110500",
            obligacion_acumulada=100,
        )
        corte_borrador = _crear_corte(
            sesion, vigencia=2026, fecha=date(2026, 3, 1), estado=EstadoCorte.BORRADOR
        )
        sesion.commit()

        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar",
            params={"corte_id_1": str(corte_registrado), "corte_id_2": str(corte_borrador)},
        )

        assert respuesta.status_code == 409
        assert respuesta.json()["detalles"]["motivo"] == "corte_no_registrado"

    def test_corte_inexistente_devuelve_404_ca07(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_a, _ = self._preparar_dos_cortes(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar",
            params={"corte_id_1": str(corte_a), "corte_id_2": str(uuid.uuid4())},
        )

        assert respuesta.status_code == 404

    def test_exportar_comparacion_devuelve_xlsx_ca08(self, cliente, sesion) -> None:
        cliente.app.dependency_overrides[obtener_usuario_actual] = lambda: _ADMINISTRADOR
        corte_a, corte_b = self._preparar_dos_cortes(sesion)

        respuesta = cliente.get(
            "/api/v1/plan-accion/comparar/exportar",
            params={"corte_id_1": str(corte_a), "corte_id_2": str(corte_b)},
        )

        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro["Comparación Plan de Acción"]
        assert hoja.cell(row=7, column=1).value == "Sector"
        valores_columna_2 = {hoja.cell(row=i, column=2).value for i in range(8, 10)}
        assert "040110500" in valores_columna_2
